from pathlib import Path

from openpyxl import load_workbook
import pytest

from data_analyser.build_template import save_template
from data_analyser.read_client import CheckFailure, validate_request
from data_analyser.registry import RegistryError, load_accountbooks
from data_analyser.snapshot import MANAGED_SHEETS


ROOT = Path(__file__).resolve().parents[1]


def test_template_builds_independently(tmp_path):
    output = save_template(tmp_path / "finance.xlsx")
    book = load_workbook(output)
    assert len(book.sheetnames) == 22
    assert "财务分析看板" in book.sheetnames
    assert len(book["财务分析看板"]._charts) == 6
    assert set(MANAGED_SHEETS) <= set(book.sheetnames)
    assert book["年度分析数据"].sheet_state == "hidden"
    book.close()
    with pytest.raises(FileExistsError):
        save_template(output)


def test_read_transport_rejects_write_methods():
    with pytest.raises(CheckFailure):
        validate_request("POST", "/jdy-fi/123/gl/v1/voucher/list", {}, {}, "123")
    validate_request("GET", "/jdy-fi/123/rpt/v1/profit", {}, None, "123")


def test_registry_rejects_session_escape(tmp_path):
    registry = tmp_path / "accountbooks.json"
    registry.write_text('{"version":2,"accountbooks":[{"key":"company_1","name":"X","company_id":"1","session_file":"../outside.json","enabled":true}]}', encoding="utf-8")
    with pytest.raises(RegistryError):
        load_accountbooks(registry)


def test_excel_macro_targets_standalone_service():
    macro = (ROOT / "excel/FinanceRefresh.bas").read_text(encoding="utf-8")
    assert 'servicePort = "18768"' in macro
    assert 'port = "18768"' in macro
    assert "\\scripts\\finance\\start-local.ps1" not in macro
    assert "\\scripts\\start-local.ps1" in macro
